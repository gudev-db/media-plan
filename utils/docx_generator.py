import io
import re
from typing import Dict, Any

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def generate_docx(params: Dict[str, Any], resultado: Dict[str, Any]) -> bytes:
    """Gera documento Word profissional a partir dos parâmetros e resultado do plano."""
    doc = Document()

    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)
    font.color.rgb = RGBColor(55, 65, 81)

    title = doc.add_heading("Plano de Mídia", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = RGBColor(31, 41, 55)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(params.get("objetivo_campanha", ""))
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(79, 70, 229)
    run.bold = True

    doc.add_paragraph()

    _add_summary_table(doc, params)
    doc.add_paragraph()
    sections = [
        ("📌 Recomendação Estratégica", resultado.get("recomendacao_estrategica", "")),
        ("📊 Distribuição de Budget", resultado.get("distribuicao_budget", "")),
        ("📈 Previsão de Resultados", resultado.get("previsao_resultados", "")),
        ("🎯 Recomendações de Público", resultado.get("recomendacoes_publico", "")),
        ("📅 Cronograma Sugerido", resultado.get("cronograma", "")),
    ]

    for title_text, body in sections:
        if body:
            heading = doc.add_heading(title_text, level=1)
            for run in heading.runs:
                run.font.color.rgb = RGBColor(31, 41, 55)
            _add_markdown_content(doc, body)
            doc.add_paragraph()
    footer_p = doc.add_paragraph()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run("Documento gerado automaticamente por IA de Planejamento de Mídia")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(156, 163, 175)
    run.italic = True

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _add_summary_table(doc: Document, params: Dict[str, Any]):
    """Adiciona tabela de resumo da campanha."""
    data = [
        ("Budget", f"R$ {params.get('budget', 0):,.2f}"),
        ("Período", params.get("periodo", "N/A")),
        ("Etapa do Funil", params.get("etapa_funil", "N/A")),
        ("Tipo de Campanha", params.get("tipo_campanha", "N/A")),
        ("Plataformas", ", ".join(params.get("ferramentas", []))),
        ("Localização Primária", params.get("localizacao_primaria", "N/A")),
        ("Tipo de Público", params.get("tipo_publico", "N/A")),
        ("Tipos de Criativo", ", ".join(params.get("tipo_criativo", []))),
    ]

    table = doc.add_table(rows=len(data), cols=2, style="Light Shading Accent 1")

    for i, (label, value) in enumerate(data):
        row = table.rows[i]
        cell_label = row.cells[0]
        cell_value = row.cells[1]
        cell_label.text = label
        cell_value.text = value

        for cell in (cell_label, cell_value):
            for paragraph in cell.paragraphs:
                paragraph.style.font.size = Pt(10)


def _add_markdown_content(doc: Document, text: str):
    """Converte markdown básico para parágrafos Word."""
    lines = text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].strip()

        if not line:
            i += 1
            continue

        # Headers markdown
        if line.startswith("###"):
            heading = doc.add_heading(line.lstrip("#").strip(), level=3)
            for run in heading.runs:
                run.font.color.rgb = RGBColor(55, 65, 81)
        elif line.startswith("##"):
            heading = doc.add_heading(line.lstrip("#").strip(), level=2)
            for run in heading.runs:
                run.font.color.rgb = RGBColor(55, 65, 81)

        # Tabelas markdown
        elif line.startswith("|") and line.endswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            _add_markdown_table(doc, table_lines)
            continue

        elif line.startswith("- ") or line.startswith("* "):
            content = line[2:].strip()
            p = doc.add_paragraph(style="List Bullet")
            _add_formatted_text(p, content)

        elif re.match(r"^\d+\.\s", line):
            content = re.sub(r"^\d+\.\s", "", line).strip()
            p = doc.add_paragraph(style="List Number")
            _add_formatted_text(p, content)
        else:
            p = doc.add_paragraph()
            _add_formatted_text(p, line)

        i += 1


def _add_markdown_table(doc: Document, table_lines: list):
    """Converte linhas de tabela markdown para tabela Word."""
    rows_data = []
    for line in table_lines:
        cells = [c.strip() for c in line.split("|") if c.strip()]
        if cells and not all(set(c).issubset({"-", ":", " "}) for c in cells):
            rows_data.append(cells)

    if not rows_data:
        return

    num_cols = max(len(row) for row in rows_data)
    table = doc.add_table(
        rows=len(rows_data), cols=num_cols, style="Light Shading Accent 1"
    )

    for i, row_data in enumerate(rows_data):
        for j, cell_text in enumerate(row_data):
            if j < num_cols:
                cell = table.rows[i].cells[j]
                cell.text = cell_text.replace("**", "")


def _add_formatted_text(paragraph, text: str):
    """Adiciona texto com formatação básica (**bold**) a um parágrafo."""
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)
